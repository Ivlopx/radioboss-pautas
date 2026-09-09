from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..core.models import ExcelSchedule, ReplacementRule, ScheduleRecord


KEY_RE = re.compile(r"\bRD[FP]\d+\b", re.IGNORECASE)
TIME_RANGE_RE = re.compile(r"\b(\d{1,2}):\d{2}(?::\d{2})?\s*-", re.IGNORECASE)
DATE_TEXT_RE = re.compile(
    r"\b(\d{1,2})\s+de\s+([a-záéíóúñ]+)\s+de\s+(\d{4})\b",
    re.IGNORECASE,
)

MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}


def _normalize(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip().casefold()


def canonical_key(value: Any) -> str:
    match = KEY_RE.search("" if value is None else str(value))
    return match.group(0).upper() if match else ""


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f"Fecha no reconocida: {value!r}")


def _hour_from_range(value: Any) -> int:
    if isinstance(value, time):
        return value.hour
    match = TIME_RANGE_RE.search("" if value is None else str(value))
    if not match:
        raise ValueError(f"Horario no reconocido: {value!r}")
    hour = int(match.group(1))
    if not 0 <= hour <= 23:
        raise ValueError(f"Hora fuera de rango: {hour}")
    return hour


def _spanish_date(text: str) -> date | None:
    match = DATE_TEXT_RE.search(text)
    if not match:
        return None
    day, month_name, year = match.groups()
    month_key = _normalize(month_name)
    month = MONTHS.get(month_key)
    if month is None:
        return None
    return date(int(year), month, int(day))


def _replacement_from_note(note: str) -> ReplacementRule | None:
    keys = [match.group(0).upper() for match in KEY_RE.finditer(note)]
    if len(keys) < 2 or "a partir" not in _normalize(note):
        return None
    effective_part = re.split(r"a\s+partir", note, flags=re.IGNORECASE, maxsplit=1)
    effective_date = _spanish_date(effective_part[-1])
    if effective_date is None:
        return None
    return ReplacementRule(keys[0], keys[-1], effective_date, note)


def read_schedule(path: str | Path) -> ExcelSchedule:
    source = Path(path)
    workbook = load_workbook(source, data_only=True, read_only=False)
    try:
        selected = None
        header_row = None
        column_map: dict[str, int] = {}
        required = {"fecha", "no. spot", "horario", "clave"}

        for sheet in workbook.worksheets:
            for row in range(1, min(sheet.max_row, 40) + 1):
                values = {_normalize(sheet.cell(row, col).value): col for col in range(1, sheet.max_column + 1)}
                if required.issubset(values):
                    selected = sheet
                    header_row = row
                    column_map = values
                    break
            if selected is not None:
                break

        if selected is None or header_row is None:
            raise ValueError("No se encontró una tabla con Fecha, No. Spot, Horario y Clave.")

        aliases = {
            "fecha": "fecha",
            "sequence": "no. spot",
            "hour": "horario",
            "key": "clave",
            "agency": "dependencia",
            "campaign": "campaña",
            "version": "versión",
            "duration": "duración",
        }

        def col(field: str) -> int | None:
            wanted = _normalize(aliases[field])
            return column_map.get(wanted)

        raw_records: list[dict[str, Any]] = []
        notes: list[str] = []
        warnings: list[str] = []
        seen_data = False

        for row in range(header_row + 1, selected.max_row + 1):
            date_value = selected.cell(row, col("fecha") or 1).value
            key_value = selected.cell(row, col("key") or 1).value
            if date_value is not None and canonical_key(key_value):
                seen_data = True
                try:
                    raw_records.append(
                        {
                            "date": _as_date(date_value),
                            "sequence": int(selected.cell(row, col("sequence") or 1).value or 0),
                            "hour": _hour_from_range(selected.cell(row, col("hour") or 1).value),
                            "key": canonical_key(key_value),
                            "agency": str(selected.cell(row, col("agency") or 1).value or "").strip(),
                            "campaign": str(selected.cell(row, col("campaign") or 1).value or "").strip(),
                            "version": str(selected.cell(row, col("version") or 1).value or "").strip(),
                            "duration": _duration(selected.cell(row, col("duration") or 1).value),
                            "row": row,
                        }
                    )
                except (TypeError, ValueError) as exc:
                    warnings.append(f"Fila {row}: {exc}")
            elif seen_data:
                values = [selected.cell(row, c).value for c in range(1, selected.max_column + 1)]
                text = " ".join(str(value).strip() for value in values if value is not None).strip()
                if len(text) >= 20 and not text.lower().startswith("tiempo total"):
                    notes.append(text)

        replacement_rules = [rule for note in notes if (rule := _replacement_from_note(note))]
        replacement_note_set = {rule.source_note for rule in replacement_rules}
        for note in notes:
            if note not in replacement_note_set and _normalize(note) != "notas":
                warnings.append(f"Nota para revisión manual: {note}")

        records: list[ScheduleRecord] = []
        for item in raw_records:
            effective_key = item["key"]
            for rule in replacement_rules:
                if effective_key == rule.old_key and item["date"] >= rule.effective_date:
                    effective_key = rule.new_key
            records.append(
                ScheduleRecord(
                    transmission_date=item["date"],
                    sequence=item["sequence"],
                    hour=item["hour"],
                    original_key=item["key"],
                    key=effective_key,
                    agency=item["agency"],
                    campaign=item["campaign"],
                    version=item["version"],
                    duration_seconds=item["duration"],
                    source_row=item["row"],
                )
            )

        if not records:
            raise ValueError("El Excel no contiene transmisiones reconocibles.")
        return ExcelSchedule(records, notes, replacement_rules, warnings, "tiempo_estado")
    finally:
        workbook.close()


def _duration(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None
