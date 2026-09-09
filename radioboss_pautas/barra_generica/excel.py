from __future__ import annotations

import re
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..tiempo_estado.excel import _normalize, canonical_key
from ..core.models import ExcelSchedule, ScheduleRecord


MONTHS = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}

SAME_MONTH_PERIOD_RE = re.compile(
    r"DEL\s+(\d{1,2})\s+AL\s+(\d{1,2})\s+DE\s+([A-ZÁÉÍÓÚÑ]+)\s+(?:DEL?\s+)?(\d{4})",
    re.IGNORECASE,
)
FULL_PERIOD_RE = re.compile(
    r"DEL\s+(\d{1,2})\s+DE\s+([A-ZÁÉÍÓÚÑ]+)\s+AL\s+"
    r"(\d{1,2})\s+DE\s+([A-ZÁÉÍÓÚÑ]+)\s+(?:DEL?\s+)?(\d{4})",
    re.IGNORECASE,
)


def read_generic_schedule(
    path: str | Path,
    program_minutes: int,
    schedule_hours: list[int],
) -> ExcelSchedule:
    if program_minutes not in (5, 10):
        raise ValueError("La duración debe ser de 5 o 10 minutos.")
    expected_hours = 2 if program_minutes == 5 else 1
    if len(schedule_hours) != expected_hours or any(not 0 <= hour <= 23 for hour in schedule_hours):
        raise ValueError(f"Selecciona {expected_hours} horario(s) válido(s).")

    workbook = load_workbook(Path(path), data_only=True, read_only=False)
    try:
        sheet = workbook.active
        start, end = _find_period(sheet)
        dates = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]

        target_text = f"estaciones con programas de {program_minutes} minutos"
        target_row = None
        for row in range(1, min(sheet.max_row, 80) + 1):
            row_text = " ".join(
                _normalize(sheet.cell(row, column).value)
                for column in range(1, min(sheet.max_column, 20) + 1)
                if sheet.cell(row, column).value is not None
            )
            if target_text in row_text:
                target_row = row
                break
        if target_row is None:
            raise ValueError(f"No se encontró la sección de programas de {program_minutes} minutos.")

        header_row = None
        for row in range(max(1, target_row - 5), target_row + 1):
            if sum(_normalize(sheet.cell(row, column).value) in WEEKDAYS for column in range(1, sheet.max_column + 1)) >= 2:
                header_row = row
        if header_row is None:
            raise ValueError("No se encontró la fila de días de la semana.")

        day_columns = [
            column
            for column in range(1, sheet.max_column + 1)
            if _normalize(sheet.cell(header_row, column).value) in WEEKDAYS
        ]
        if len(day_columns) != len(dates):
            raise ValueError(
                f"El periodo contiene {len(dates)} fechas, pero se encontraron {len(day_columns)} columnas de día."
            )

        source_rows = [target_row]
        if program_minutes == 5:
            second_row = target_row + 1
            if "horario 2" not in _normalize(sheet.cell(second_row, 2).value):
                raise ValueError("No se encontró la fila HORARIO 2.")
            source_rows.append(second_row)

        records: list[ScheduleRecord] = []
        warnings: list[str] = []
        sequence = 1
        for schedule_index, source_row in enumerate(source_rows):
            for transmission_date, column in zip(dates, day_columns):
                raw = sheet.cell(source_row, column).value
                key = canonical_key(raw)
                if not key:
                    warnings.append(f"Celda {sheet.cell(source_row, column).coordinate}: no contiene una clave RDP.")
                    continue
                parts = [part.strip() for part in str(raw).split("/")]
                records.append(
                    ScheduleRecord(
                        transmission_date=transmission_date,
                        sequence=sequence,
                        hour=schedule_hours[schedule_index],
                        original_key=key,
                        key=key,
                        agency=parts[1] if len(parts) > 1 else "",
                        campaign=parts[2] if len(parts) > 2 else "Barra genérica nacional",
                        version=" / ".join(parts[3:]) if len(parts) > 3 else "",
                        duration_seconds=program_minutes * 60,
                        source_row=source_row,
                    )
                )
                sequence += 1

        for row in range(max(source_rows) + 1, min(sheet.max_row, max(source_rows) + 12) + 1):
            text = " ".join(
                str(sheet.cell(row, column).value).strip()
                for column in range(1, sheet.max_column + 1)
                if sheet.cell(row, column).value not in (None, ".")
            ).strip()
            if len(text) >= 20 and not canonical_key(text):
                warnings.append(f"Nota para revisión manual: {text}")

        if not records:
            raise ValueError("La sección elegida no contiene programas reconocibles.")
        return ExcelSchedule(records, [], [], warnings, "barra_generica")
    finally:
        workbook.close()


WEEKDAYS = {"lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo"}


def _find_period(sheet: Any) -> tuple[date, date]:
    for row in range(1, min(sheet.max_row, 30) + 1):
        for column in range(1, min(sheet.max_column, 20) + 1):
            value = sheet.cell(row, column).value
            if not isinstance(value, str):
                continue
            full = FULL_PERIOD_RE.search(value)
            if full:
                start_day, start_month_name, end_day, end_month_name, year = full.groups()
                start_month = MONTHS[_normalize(start_month_name)]
                end_month = MONTHS[_normalize(end_month_name)]
                start = date(int(year), start_month, int(start_day))
                end_year = int(year) + (1 if end_month < start_month else 0)
                return start, date(end_year, end_month, int(end_day))
            same = SAME_MONTH_PERIOD_RE.search(value)
            if same:
                start_day, end_day, month_name, year = same.groups()
                month = MONTHS[_normalize(month_name)]
                return date(int(year), month, int(start_day)), date(int(year), month, int(end_day))
    raise ValueError("No se pudo determinar el periodo de la barra genérica.")
