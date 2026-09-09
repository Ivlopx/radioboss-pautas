from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from radioboss_pautas.barra_generica.excel import read_generic_schedule
from radioboss_pautas.core.ini import date_to_delphi_serial, read_ini, render_ini
from radioboss_pautas.core.planner import build_plan
from radioboss_pautas.service import AnalysisResult, analyze, generate, restore_backup
from radioboss_pautas.tiempo_estado.excel import read_schedule


@pytest.fixture
def sample_files(tmp_path: Path) -> tuple[Path, Path, Path]:
    ini = tmp_path / "adscheduler.ini.txt"
    time_state = tmp_path / "tiempo_estado.xlsx"
    generic = tmp_path / "barra_generica.xlsx"
    _write_sample_ini(ini)
    _write_time_state_workbook(time_state)
    _write_generic_workbook(generic)
    return ini, time_state, generic


def _write_sample_ini(path: Path) -> None:
    hours = "".join("1" if 6 <= value <= 21 else "0" for value in range(24))
    minutes = "".join("1" if value in {5, 25, 45} else "0" for value in range(60))
    tracks: list[str] = []
    track_specs = [(9, 5), (10, 25), (11, 45), (12, 5), (13, 25)]
    for index, (hour, minute) in enumerate(track_specs):
        lines = ["    item", "      Value.Marks = <"]
        for weekday in range(1, 8):
            lines.extend(
                [
                    "        item",
                    f"          Hour = {hour}",
                    f"          Minute = {minute}",
                    f"          Day = {weekday}",
                    "        end",
                ]
            )
        lines[-1] += ">"
        lines.extend(
            [
                "      Value.Priority = 49",
                "      Value.Active = True",
                "      Value.IsStart = False",
                "      Value.IsEnd = False",
                f"      Value.DisplayName = 'MANUAL_{index + 1}'",
                "      Value.AltTracks.Strings = (",
                f"        'C:\\RadioBOSS\\manual_{index + 1}.mp3')",
                "    end>" if index == len(track_specs) - 1 else "    end",
            ]
        )
        tracks.append("\n".join(lines))

    content = "\n".join(
        [
            "object TAdsSchedulerIni",
            "  Version = 1",
            f"  SHours = '{hours}'",
            f"  SMinutes = '{minutes}'",
            "  Second = 0",
            "  AdTracks = <",
            *tracks,
            "  WorkFolder = 'C:\\RadioBOSS\\AdPlaylists\\'",
            "end",
            "",
        ]
    )
    path.write_bytes(content.encode("ascii"))


def _write_time_state_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pauta"
    sheet.append(["Documento de prueba sin datos operativos"])
    for _ in range(4):
        sheet.append([])
    sheet.append(["Fecha", "No. Spot", "Horario", "Clave", "Dependencia", "Campaña", "Versión", "Duración"])

    start = date(2026, 9, 7)
    hours = list(range(6, 24)) + list(range(6, 11))
    keys = [f"RDF{200 + value}2026" for value in range(15)]
    old_keys = {3: "RDF2382026", 9: "RDF2392026"}
    for day_offset in range(8):
        transmission_date = start + timedelta(days=day_offset)
        for sequence, hour in enumerate(hours, start=1):
            key = keys[(sequence + day_offset) % len(keys)]
            if day_offset == 7 and sequence in old_keys:
                key = old_keys[sequence]
            sheet.append(
                [
                    transmission_date,
                    sequence,
                    f"{hour:02d}:00:00 - {(hour + 1) % 24:02d}:00:00",
                    key,
                    "Dependencia de prueba",
                    "Campaña de prueba",
                    "Versión única",
                    30,
                ]
            )

    sheet.append(["NOTAS"])
    sheet.append(["La campaña RDF2382026 termina; favor de sustituirla por RDF2412026 a partir del 14 de septiembre de 2026."])
    sheet.append(["La campaña RDF2392026 termina; favor de sustituirla por RDF2422026 a partir del 14 de septiembre de 2026."])
    workbook.save(path)


def _write_generic_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GENERICA"
    sheet["A2"] = "PROGRAMAS CORRESPONDIENTES AL PERIODO DEL 07 AL 14 DE SEPTIEMBRE DEL 2026"
    weekdays = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO", "LUNES"]
    for column, weekday in enumerate(weekdays, start=3):
        sheet.cell(5, column, weekday)

    sheet["A6"] = "ESTACIONES CON PROGRAMAS DE 5 MINUTOS"
    sheet["B6"] = "HORARIO 1"
    sheet["B7"] = "HORARIO 2"
    five_keys = [f"RDP{100 + value}2026" for value in range(6)]
    for offset, column in enumerate(range(3, 11)):
        sheet.cell(6, column, f"{five_keys[offset % 6]} / Institución / Programa")
        sheet.cell(7, column, f"{five_keys[(offset + 2) % 6]} / Institución / Programa")

    sheet["A8"] = "ESTACIONES CON PROGRAMAS DE 10 MINUTOS"
    sheet["B8"] = "HORARIO 1"
    ten_keys = [f"RDP{200 + value}2026" for value in range(3)]
    for offset, column in enumerate(range(3, 11)):
        sheet.cell(8, column, f"{ten_keys[offset % 3]} / Institución / Programa")
    sheet["A9"] = "Transmitir el Himno Nacional conforme a las indicaciones aplicables."
    workbook.save(path)


def test_reads_excel_and_applies_note_replacements(sample_files: tuple[Path, Path, Path]) -> None:
    _, excel, _ = sample_files
    schedule = read_schedule(excel)
    assert len(schedule.records) == 184
    assert len({record.transmission_date for record in schedule.records}) == 8
    assert len(schedule.replacement_rules) == 2
    september_14 = [record for record in schedule.records if record.transmission_date == date(2026, 9, 14)]
    assert any(record.original_key == "RDF2392026" and record.key == "RDF2422026" for record in september_14)
    assert any(record.original_key == "RDF2382026" and record.key == "RDF2412026" for record in september_14)


def test_reads_template_masks_and_tracks(sample_files: tuple[Path, Path, Path]) -> None:
    ini, _, _ = sample_files
    template = read_ini(ini)
    assert template.configured_hours == list(range(6, 22))
    assert template.configured_minutes == [5, 25, 45]
    assert len(template.existing_tracks) == 5


def test_reads_both_generic_program_sections(sample_files: tuple[Path, Path, Path]) -> None:
    _, _, generic_excel = sample_files
    five = read_generic_schedule(generic_excel, 5, [9, 9])
    ten = read_generic_schedule(generic_excel, 10, [12])
    assert len(five.records) == 16
    assert len({record.key for record in five.records}) == 6
    assert {record.hour for record in five.records} == {9}
    assert len(ten.records) == 8
    assert len({record.key for record in ten.records}) == 3
    assert {record.hour for record in ten.records} == {12}
    assert min(record.transmission_date for record in five.records) == date(2026, 9, 7)
    assert max(record.transmission_date for record in five.records) == date(2026, 9, 14)
    assert all("Himno Nacional" in warning for warning in five.warnings)


def test_generic_programs_choose_least_loaded_ini_block(sample_files: tuple[Path, Path, Path], tmp_path: Path) -> None:
    ini, _, generic_excel = sample_files
    schedule = read_generic_schedule(generic_excel, 5, [9, 9])
    template = read_ini(ini)
    for key in {record.key for record in schedule.records}:
        (tmp_path / f"{key} programa.mp3").touch()
    plan = build_plan(schedule, template, tmp_path, strategy="least_loaded")
    monday = [item for item in plan.assignments if item.transmission_date == date(2026, 9, 7)]
    assert [item.minute for item in monday] == [25, 45]
    assert plan.managed_prefix == "AUTO_BGN_"
    assert plan.display_group == "BARRA GENERICA NACIONAL"

    rendered = render_ini(template, plan.assignments, plan.enabled_hours_added, plan.managed_prefix, plan.display_group)
    assert rendered.count("Value.DisplayName = 'AUTO_BGN_") > 0
    assert "Value.DisplayName = 'AUTO_TE_" not in rendered


def test_generic_section_rejects_an_hour_disabled_in_ini(sample_files: tuple[Path, Path, Path], tmp_path: Path) -> None:
    ini, _, generic_excel = sample_files
    schedule = read_generic_schedule(generic_excel, 10, [9])
    for key in {record.key for record in schedule.records}:
        (tmp_path / f"{key} programa.mp3").touch()
    with pytest.raises(ValueError, match="23:00"):
        analyze(generic_excel, tmp_path, ini, mode="barra_generica", program_minutes=10, schedule_hours=[23])


def test_regenerating_one_section_preserves_the_other(sample_files: tuple[Path, Path, Path], tmp_path: Path) -> None:
    ini, excel, generic_excel = sample_files
    template = read_ini(ini)
    time_state = read_schedule(excel)
    generic = read_generic_schedule(generic_excel, 10, [12])
    audio_folder = tmp_path / "audios"
    audio_folder.mkdir()
    for key in {record.key for record in time_state.records + generic.records}:
        (audio_folder / f"{key} audio.mp3").touch()

    te_plan = build_plan(time_state, template, audio_folder, strategy="active_only")
    with_te = render_ini(template, te_plan.assignments, te_plan.enabled_hours_added, te_plan.managed_prefix, te_plan.display_group)
    combined_path = tmp_path / "combined.ini.txt"
    combined_path.write_bytes(with_te.encode("ascii"))

    combined_template = read_ini(combined_path)
    generic_plan = build_plan(generic, combined_template, audio_folder, strategy="least_loaded")
    combined = render_ini(combined_template, generic_plan.assignments, generic_plan.enabled_hours_added, generic_plan.managed_prefix, generic_plan.display_group)
    assert "Value.DisplayName = 'AUTO_TE_" in combined
    assert "Value.DisplayName = 'AUTO_BGN_" in combined


def test_plan_and_render_with_matching_audio_files(sample_files: tuple[Path, Path, Path], tmp_path: Path) -> None:
    ini, excel, _ = sample_files
    schedule = read_schedule(excel)
    template = read_ini(ini)
    for key in {record.key for record in schedule.records}:
        (tmp_path / f"{key} material.mp3").touch()

    plan = build_plan(schedule, template, tmp_path)
    assert plan.can_generate
    assert len(plan.assignments) == 184
    assert plan.enabled_hours_added == [22, 23]
    assert all(assignment.minute in template.configured_minutes for assignment in plan.assignments)
    first_six = next(assignment for assignment in plan.assignments if assignment.transmission_date == date(2026, 9, 7) and assignment.hour == 6)
    assert first_six.minute == template.configured_minutes[0]
    assert first_six.block_origin == "Creado"

    rendered = render_ini(template, plan.assignments, plan.enabled_hours_added)
    assert "Value.DisplayName = 'AUTO_TE_" in rendered
    assert f"Value.Start = {date_to_delphi_serial(date(2026, 9, 7))}" in rendered
    assert "SHours = '000000111111111111111111'" in rendered
    assert rendered.count("object TAdsSchedulerIni") == 1
    assert rendered.rstrip().endswith("end")

    generated_path = tmp_path / "generated.ini.txt"
    generated_path.write_bytes(rendered.encode("ascii"))
    regenerated_template = read_ini(generated_path)
    regenerated = render_ini(regenerated_template, plan.assignments, plan.enabled_hours_added)
    assert regenerated.count("Value.DisplayName = 'AUTO_TE_") == rendered.count("Value.DisplayName = 'AUTO_TE_")


def test_generation_creates_backup_and_backup_can_be_restored(sample_files: tuple[Path, Path, Path], tmp_path: Path) -> None:
    ini, excel_path, _ = sample_files
    excel = read_schedule(excel_path)
    template = read_ini(ini)
    audio_folder = tmp_path / "Material Radio"
    audio_folder.mkdir()
    for key in {record.key for record in excel.records}:
        (audio_folder / f"{key} audio.mp3").touch()
    plan = build_plan(excel, template, audio_folder)
    result = generate(AnalysisResult(excel, template, plan))

    assert result.output_path == ini.resolve()
    assert result.backup_path is not None and result.backup_path.exists()
    assert b"AUTO_TE_" in ini.read_bytes()
    assert b"AUTO_TE_" not in result.backup_path.read_bytes()

    restore_backup(result.backup_path, ini)
    assert b"AUTO_TE_" not in ini.read_bytes()
